"""Export stratified conversation examples from cleaned ACT data."""

## IMPORTS ##
# Standard library
import logging
import random
import re
from pathlib import Path

# External
import polars as pl
from docx import Document
## _______ ##


## HELPER FUNCTIONS ##
def _add_topics(
    df: pl.DataFrame,
    topic_contact_reasons: dict[str, list[str]],
    logger: logging.Logger,
) -> pl.DataFrame:
    logger.info("Adding topics to conversations.")

    contact_reason_to_topic = {}

    for topic, contact_reasons in topic_contact_reasons.items():
        for contact_reason in contact_reasons:
            contact_reason_to_topic[contact_reason] = topic

    rows = []

    for row in df.iter_rows(named=True):
        contact_reasons = row["contact_reasons"]

        if contact_reasons is None:
            logger.warning(
                "Conversation skipped because contact_reasons is None."
            )
            continue

        topics = []

        for contact_reason in contact_reasons:
            if contact_reason in contact_reason_to_topic:
                topics.append(
                    contact_reason_to_topic[contact_reason]
                )
            else:
                logger.warning(
                    f"Unknown contact reason encountered: "
                    f"{contact_reason}"
                )

        topics = sorted(set(topics))

        if not topics:
            logger.warning(
                f"Conversation contained no valid topics. "
                f"Contact reasons: {contact_reasons}"
            )
            continue

        row["topics"] = topics
        rows.append(row)

    logger.info(
        f"Kept {len(rows)} conversations after topic assignment."
    )

    return pl.DataFrame(rows)


def _sanitize_folder_name(value: str) -> str:
    value = value.strip().title()
    value = re.sub(r"[^a-zA-Z0-9æøåÆØÅ_-]+", "_", value)
    value = re.sub(r"_+", "_", value)

    return value.strip("_")


def _calculate_topic_counts(
    rows: list[dict],
    topic_contact_reasons: dict[str, list[str]],
) -> dict[str, int]:
    topic_counts = {
        topic: 0
        for topic in topic_contact_reasons.keys()
    }

    for row in rows:
        for topic in row["topics"]:
            topic_counts[topic] += 1

    return topic_counts


def _calculate_topic_imbalance_score(
    rows: list[dict],
    topic_contact_reasons: dict[str, list[str]],
) -> int:
    topic_counts = _calculate_topic_counts(
        rows=rows,
        topic_contact_reasons=topic_contact_reasons,
    )

    counts = list(topic_counts.values())

    return max(counts) - min(counts)


def _calculate_average_wordcount(
    rows: list[dict],
) -> float:
    return sum(row["text_words"] for row in rows) / len(rows)


def _calculate_wordcount_distance_from_target(
    rows: list[dict],
    target: int,
) -> float:
    avg_wordcount = _calculate_average_wordcount(rows)

    return abs(avg_wordcount - target)


def _is_average_wordcount_valid(
    rows: list[dict],
    target: int,
    divagation: int,
) -> bool:
    avg_wordcount = _calculate_average_wordcount(rows)

    return target - divagation <= avg_wordcount <= target + divagation


def _find_most_underrepresented_topics(
    selected_rows: list[dict],
    topic_contact_reasons: dict[str, list[str]],
) -> list[str]:
    topic_counts = _calculate_topic_counts(
        rows=selected_rows,
        topic_contact_reasons=topic_contact_reasons,
    )

    lowest_topic_count = min(topic_counts.values())

    return [
        topic
        for topic, count in topic_counts.items()
        if count == lowest_topic_count
    ]


def _row_contains_any_topic(
    row: dict,
    topics: list[str],
) -> bool:
    return any(
        topic in row["topics"]
        for topic in topics
    )


def _find_best_candidate_rows(
    selected_rows: list[dict],
    candidate_rows: list[dict],
    topic_contact_reasons: dict[str, list[str]],
) -> list[dict]:
    best_rows = []
    best_score = None

    for row in candidate_rows:
        score = _calculate_topic_imbalance_score(
            rows=[*selected_rows, row],
            topic_contact_reasons=topic_contact_reasons,
        )

        if best_score is None or score < best_score:
            best_score = score
            best_rows = [row]

        elif score == best_score:
            best_rows.append(row)

    return best_rows


def _select_topic_guided_sample_once(
    candidate_rows: list[dict],
    n_conversations: int,
    topic_contact_reasons: dict[str, list[str]],
    rng: random.Random,
    logger: logging.Logger,
) -> list[dict]:
    selected_rows = []
    remaining_rows = candidate_rows.copy()

    rng.shuffle(remaining_rows)

    while (
        len(selected_rows) < n_conversations
        and remaining_rows
    ):
        most_underrepresented_topics = _find_most_underrepresented_topics(
            selected_rows=selected_rows,
            topic_contact_reasons=topic_contact_reasons,
        )

        topic_relevant_rows = [
            row
            for row in remaining_rows
            if _row_contains_any_topic(
                row=row,
                topics=most_underrepresented_topics,
            )
        ]

        if not topic_relevant_rows:
            logger.debug(
                f"No remaining rows found for currently underrepresented "
                f"topics: {most_underrepresented_topics}. "
                f"Using all remaining rows instead."
            )

            topic_relevant_rows = remaining_rows

        best_candidate_rows = _find_best_candidate_rows(
            selected_rows=selected_rows,
            candidate_rows=topic_relevant_rows,
            topic_contact_reasons=topic_contact_reasons,
        )

        selected_row = rng.choice(best_candidate_rows)

        selected_rows.append(selected_row)
        remaining_rows.remove(selected_row)

    return selected_rows


def _select_best_topic_guided_sample(
    candidate_rows: list[dict],
    n_conversations: int,
    target: int,
    divagation: int,
    topic_contact_reasons: dict[str, list[str]],
    logger: logging.Logger,
    max_attempts: int,
    max_sampling_rounds: int,
    rng: random.Random,
) -> list[dict]:
    best_invalid_sample = None
    best_invalid_score = None

    for sampling_round in range(1, max_sampling_rounds + 1):
        logger.info(
            f"Starting sampling round {sampling_round} of "
            f"{max_sampling_rounds}."
        )

        best_valid_sample = None
        best_valid_score = None

        for attempt in range(1, max_attempts + 1):
            selected_rows = _select_topic_guided_sample_once(
                candidate_rows=candidate_rows,
                n_conversations=n_conversations,
                topic_contact_reasons=topic_contact_reasons,
                rng=rng,
                logger=logger,
            )

            if len(selected_rows) != n_conversations:
                raise RuntimeError(
                    f"Sampling returned {len(selected_rows)} conversations "
                    f"although {n_conversations} were requested."
                )

            avg_wordcount = _calculate_average_wordcount(selected_rows)
            wordcount_distance = _calculate_wordcount_distance_from_target(
                rows=selected_rows,
                target=target,
            )
            topic_imbalance_score = _calculate_topic_imbalance_score(
                rows=selected_rows,
                topic_contact_reasons=topic_contact_reasons,
            )

            invalid_score = (
                wordcount_distance,
                topic_imbalance_score,
            )

            if (
                best_invalid_score is None
                or invalid_score < best_invalid_score
            ):
                best_invalid_sample = selected_rows
                best_invalid_score = invalid_score

            if not _is_average_wordcount_valid(
                rows=selected_rows,
                target=target,
                divagation=divagation,
            ):
                continue

            valid_score = (
                topic_imbalance_score,
                wordcount_distance,
            )

            if (
                best_valid_score is None
                or valid_score < best_valid_score
            ):
                best_valid_sample = selected_rows
                best_valid_score = valid_score

        if best_valid_sample is not None:
            logger.info(
                f"Valid sample found in sampling round {sampling_round}. "
                f"Average wordcount="
                f"{_calculate_average_wordcount(best_valid_sample):.1f}, "
                f"topic_imbalance_score="
                f"{_calculate_topic_imbalance_score(
                    rows=best_valid_sample,
                    topic_contact_reasons=topic_contact_reasons,
                )}."
            )

            return best_valid_sample

        logger.warning(
            f"No valid sample found in sampling round {sampling_round}. "
            f"Trying another round if available."
        )

    if best_invalid_sample is None:
        logger.warning(
            "No valid or invalid sample could be selected. "
            "Returning an empty list."
        )
        return []

    logger.warning(
        f"No valid sample found after {max_sampling_rounds} sampling rounds. "
        f"Returning the closest invalid sample. "
        f"Average wordcount="
        f"{_calculate_average_wordcount(best_invalid_sample):.1f}, "
        f"target={target}, divagation={divagation}, "
        f"topic_imbalance_score="
        f"{_calculate_topic_imbalance_score(
            rows=best_invalid_sample,
            topic_contact_reasons=topic_contact_reasons,
        )}."
    )

    return best_invalid_sample


def _select_balanced_rows(
    df: pl.DataFrame,
    gender: str,
    age_group: str,
    n_conversations: int,
    lower_limit: int,
    upper_limit: int,
    target: int,
    divagation: int,
    topic_contact_reasons: dict[str, list[str]],
    logger: logging.Logger,
    max_attempts: int,
    max_sampling_rounds: int,
    rng: random.Random,
) -> list[dict]:
    logger.info(
        f"Selecting {n_conversations} conversations for "
        f"gender={gender}, age_group={age_group}."
    )

    candidate_rows = (
        df
        .filter(
            (pl.col("gender") == gender)
            & (pl.col("age_group") == age_group)
            & (pl.col("text_words") >= lower_limit)
            & (pl.col("text_words") <= upper_limit)
        )
        .to_dicts()
    )

    logger.info(
        f"Found {len(candidate_rows)} candidate rows for "
        f"gender={gender}, age_group={age_group}."
    )

    if not candidate_rows:
        logger.warning(
            f"No eligible conversations found for gender={gender}, "
            f"age_group={age_group}."
        )
        return []

    if len(candidate_rows) <= n_conversations:
        if len(candidate_rows) < n_conversations:
            logger.warning(
                f"Only {len(candidate_rows)} eligible conversations found for "
                f"gender={gender}, age_group={age_group}. "
                f"Requested {n_conversations}. "
                f"Exporting all available conversations without further balancing."
            )
        else:
            logger.info(
                f"Exactly {n_conversations} eligible conversations found for "
                f"gender={gender}, age_group={age_group}. "
                f"No sampling required."
            )

        return candidate_rows

    return _select_best_topic_guided_sample(
        candidate_rows=candidate_rows,
        n_conversations=n_conversations,
        target=target,
        divagation=divagation,
        topic_contact_reasons=topic_contact_reasons,
        logger=logger,
        max_attempts=max_attempts,
        max_sampling_rounds=max_sampling_rounds,
        rng=rng,
    )


def _save_conversation_to_docx(
    row: dict,
    output_path: Path,
    logger: logging.Logger,
) -> None:
    logger.info(f"Saving conversation docx to {output_path}")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    doc.add_paragraph(f"Køn : {row['requester_gender_name']}")
    doc.add_paragraph(f"Alder : {row['requester_age_name']}")
    doc.add_paragraph(f"Henvendelsesårsager : {row['contact_reasons']}")
    doc.add_paragraph(f"Tekst : {row['msg_tekst']}")

    doc.save(output_path)


def _save_selected_conversations(
    selected_rows: list[dict],
    output_dir_data_folder: str,
    gender: str,
    age_group: str,
    logger: logging.Logger,
) -> None:
    output_folder = (
        Path(output_dir_data_folder)
        / _sanitize_folder_name(gender)
        / _sanitize_folder_name(age_group)
    )

    for i, row in enumerate(selected_rows, start=1):
        output_path = output_folder / f"conversation_{i:03d}.docx"

        _save_conversation_to_docx(
            row=row,
            output_path=output_path,
            logger=logger,
        )


def _log_selected_distribution(
    selected_rows: list[dict],
    topic_contact_reasons: dict[str, list[str]],
    gender: str,
    age_group: str,
    logger: logging.Logger,
) -> None:
    topic_counts = _calculate_topic_counts(
        rows=selected_rows,
        topic_contact_reasons=topic_contact_reasons,
    )

    logger.info(
        f"Topic distribution for gender={gender}, age_group={age_group}: "
        f"{topic_counts}"
    )

    logger.info(
        f"Average wordcount for gender={gender}, age_group={age_group}: "
        f"{_calculate_average_wordcount(selected_rows):.1f}"
    )


def _build_contact_reason_to_topic(
    topic_contact_reasons: dict[str, list[str]],
) -> dict[str, str]:
    return {
        contact_reason: topic
        for topic, contact_reasons in topic_contact_reasons.items()
        for contact_reason in contact_reasons
    }


def _calculate_unique_contact_reasons(
    rows: list[dict],
) -> set[str]:
    return {
        contact_reason
        for row in rows
        for contact_reason in row["contact_reasons"]
    }


def _calculate_unique_contact_reasons_by_topic(
    rows: list[dict],
    topic_contact_reasons: dict[str, list[str]],
) -> dict[str, set[str]]:
    contact_reason_to_topic = _build_contact_reason_to_topic(
        topic_contact_reasons
    )

    unique_by_topic = {
        topic: set()
        for topic in topic_contact_reasons.keys()
    }

    for row in rows:
        for contact_reason in row["contact_reasons"]:
            topic = contact_reason_to_topic.get(contact_reason)

            if topic is not None:
                unique_by_topic[topic].add(contact_reason)

    return unique_by_topic


def _format_export_summary(
    selected_rows_by_group: dict[tuple[str, str], list[dict]],
    topic_contact_reasons: dict[str, list[str]],
) -> str:
    lines = [
        "Stratified conversation export summary",
        "=" * 40,
        "",
    ]

    for (gender, age_group), selected_rows in selected_rows_by_group.items():
        topic_counts = _calculate_topic_counts(
            rows=selected_rows,
            topic_contact_reasons=topic_contact_reasons,
        )
        unique_contact_reasons = _calculate_unique_contact_reasons(
            selected_rows
        )
        unique_contact_reasons_by_topic = (
            _calculate_unique_contact_reasons_by_topic(
                rows=selected_rows,
                topic_contact_reasons=topic_contact_reasons,
            )
        )

        lines.extend(
            [
                f"Gender: {gender}",
                f"Age group: {age_group}",
                f"Total conversations: {len(selected_rows)}",
                f"Average wordcount: "
                f"{_calculate_average_wordcount(selected_rows):.1f}",
                "",
                "Topic distribution:",
            ]
        )

        for topic, count in topic_counts.items():
            lines.append(f"- {topic}: {count}")

        lines.extend(
            [
                "",
                f"Unique contact_reasons: "
                f"{len(unique_contact_reasons)}",
                "",
                "Unique contact_reasons by topic:",
            ]
        )

        for topic, contact_reasons in unique_contact_reasons_by_topic.items():
            lines.append(
                f"- {topic}: {len(contact_reasons)}"
            )

        lines.append("")
        lines.append("-" * 40)
        lines.append("")

    return "\n".join(lines)


def _save_export_summary(
    summary_text: str,
    output_dir_data_folder: str,
    logger: logging.Logger,
) -> None:
    output_path = Path(output_dir_data_folder) / "export_summary.txt"

    logger.info(f"Saving export summary to {output_path}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(summary_text, encoding="utf-8")


## MAIN FUNCTIONALITY ##
def export_stratified_conversations(
    input_dir_data: str,
    output_dir_data_folder: str,
    logger: logging.Logger,
    topic_contact_reasons: dict[str, list[str]],
    sample_plan: dict[str, dict[str, int]],
    wordcount_settings: dict[str, dict[str, int]],
    divagation: int = 150,
    max_attempts: int = 1000,
    max_sampling_rounds: int = 3,
    random_seed: int = 42,
) -> None:
    logger.info(f"Reading input data from {input_dir_data}")

    df = pl.read_parquet(input_dir_data)

    logger.info(f"Loaded dataset with {df.height} rows and {df.width} columns.")

    df = _add_topics(
        df=df,
        topic_contact_reasons=topic_contact_reasons,
        logger=logger,
    )

    rng = random.Random(random_seed)

    selected_rows_by_group = {}

    for gender, age_group_plan in sample_plan.items():
        gender_wordcount_settings = wordcount_settings[gender]

        for age_group, n_conversations in age_group_plan.items():
            selected_rows = _select_balanced_rows(
                df=df,
                gender=gender,
                age_group=age_group,
                n_conversations=n_conversations,
                lower_limit=gender_wordcount_settings["lower_limit"],
                upper_limit=gender_wordcount_settings["upper_limit"],
                target=gender_wordcount_settings["target"],
                divagation=divagation,
                topic_contact_reasons=topic_contact_reasons,
                logger=logger,
                max_attempts=max_attempts,
                max_sampling_rounds=max_sampling_rounds,
                rng=rng,
            )

            if not selected_rows:
                logger.warning(
                    f"No conversations exported for gender={gender}, "
                    f"age_group={age_group}."
                )
                continue

            selected_rows_by_group[(gender, age_group)] = selected_rows

            _log_selected_distribution(
                selected_rows=selected_rows,
                topic_contact_reasons=topic_contact_reasons,
                gender=gender,
                age_group=age_group,
                logger=logger,
            )

            _save_selected_conversations(
                selected_rows=selected_rows,
                output_dir_data_folder=output_dir_data_folder,
                gender=gender,
                age_group=age_group,
                logger=logger,
            )

    if selected_rows_by_group:
        summary_text = _format_export_summary(
            selected_rows_by_group=selected_rows_by_group,
            topic_contact_reasons=topic_contact_reasons,
        )

        logger.info("\n" + summary_text)

        _save_export_summary(
            summary_text=summary_text,
            output_dir_data_folder=output_dir_data_folder,
            logger=logger,
        )

    logger.info("Finished exporting stratified conversation examples.")
## __________________ ##