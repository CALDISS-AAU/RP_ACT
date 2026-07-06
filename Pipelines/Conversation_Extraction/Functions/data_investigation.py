"""Generate descriptive statistics and reports for cleaned ACT conversation data.
   The module calculates wordcount, topic, and contact-reason statistics,
       creates distribution charts, and exports the results as a Word
       report for exploratory data analysis.
"""

## IMPORTS ##
# Standard library
import logging
from math import ceil
from pathlib import Path

# External
import matplotlib.pyplot as plt
import polars as pl
from docx import Document
from docx.shared import Inches
## _______ ##


## HELPER FUNCTIONS ##
def _shortest_after_removing_bottom_percent(
    df: pl.DataFrame,
    bottom_percent: float,
    logger: logging.Logger,
) -> int | None:
    """Return the shortest wordcount after removing the lowest percentage of texts."""
    values = (
        df
        .select(pl.col("text_words").drop_nulls().sort())
        .to_series()
        .to_list()
    )

    if not values:
        logger.warning("No wordcount values found.")
        return None

    n_to_remove = ceil(len(values) * (bottom_percent / 100))

    if n_to_remove >= len(values):
        logger.warning(
            f"Bottom percent removal removes entire dataset: "
            f"n_values={len(values)}, "
            f"n_to_remove={n_to_remove}, "
            f"bottom_percent={bottom_percent}"
        )
        return None

    return values[n_to_remove]


def _longest_after_removing_top_percent(
    df: pl.DataFrame,
    top_percent: float,
    logger: logging.Logger,
) -> int | None:
    """Return the longest wordcount after removing the highest percentage of texts."""
    values = (
        df
        .select(pl.col("text_words").drop_nulls().sort())
        .to_series()
        .to_list()
    )

    if not values:
        logger.warning("No wordcount values found.")
        return None

    n_to_remove = ceil(len(values) * (top_percent / 100))

    if n_to_remove >= len(values):
        logger.warning(
            f"Top percent removal removes entire dataset: "
            f"n_values={len(values)}, "
            f"n_to_remove={n_to_remove}, "
            f"top_percent={top_percent}"
        )
        return None

    return values[-(n_to_remove + 1)]


def _calculate_wordcount_stats_for_df(
    df: pl.DataFrame,
    bottom_percent: float,
    logger: logging.Logger,
) -> dict:
    """Calculate summary wordcount statistics for a dataframe."""
    logger.debug(f"Calculating wordcount stats for {df.height} rows.")

    shortest_after_trim = _shortest_after_removing_bottom_percent(
        df=df,
        bottom_percent=bottom_percent,
        logger=logger,
    )

    longest_after_trim = _longest_after_removing_top_percent(
        df=df,
        top_percent=bottom_percent,
        logger=logger,
    )

    stats = (
        df
        .select(
            pl.len().alias("n_texts"),
            pl.col("text_words").min().alias("min_wordcount"),
            pl.col("text_words").max().alias("max_wordcount"),
            pl.col("text_words").mean().round(1).alias("avg_wordcount"),
            pl.col("text_words").median().alias("med_wordcount"),
        )
        .to_dicts()[0]
    )

    stats[f"shortest_after_removing_bottom_{bottom_percent:g}_percent"] = (
        shortest_after_trim
    )

    stats[f"longest_after_removing_top_{bottom_percent:g}_percent"] = (
        longest_after_trim
    )

    return stats


def _calculate_grouped_wordcount_stats(
    df: pl.DataFrame,
    group_cols: list[str] | None,
    bottom_percent: float,
    logger: logging.Logger,
) -> pl.DataFrame:
    """Calculate wordcount statistics grouped by one or more columns."""
    if group_cols is None:
        logger.info("Calculating wordcount stats for entire dataset.")

        stats = _calculate_wordcount_stats_for_df(
            df=df,
            bottom_percent=bottom_percent,
            logger=logger,
        )

        return pl.DataFrame([{"group": "Entire dataset", **stats}])

    logger.info(f"Calculating wordcount stats grouped by: {', '.join(group_cols)}")

    df = df.filter(
        pl.all_horizontal([
            pl.col(col).is_not_null()
            for col in group_cols
        ])
    )

    rows = []

    for group_key, group_df in df.group_by(group_cols, maintain_order=True):
        if not isinstance(group_key, tuple):
            group_key = (group_key,)

        group_values = dict(zip(group_cols, group_key))

        stats = _calculate_wordcount_stats_for_df(
            df=group_df,
            bottom_percent=bottom_percent,
            logger=logger,
        )

        rows.append({**group_values, **stats})

    return pl.DataFrame(rows)


def _create_wordcount_report_sections(
    df: pl.DataFrame,
    bottom_percent: float,
    logger: logging.Logger,
) -> dict[str, dict]:
    """Create report sections for wordcount statistics."""
    logger.info("Creating wordcount report sections.")

    sections = {
        "wordcount_entire_dataset": {
            "title": "Wordcount: Entire dataset",
            "source_df": df,
            "stats_df": _calculate_grouped_wordcount_stats(
                df=df,
                group_cols=None,
                bottom_percent=bottom_percent,
                logger=logger,
            ),
            "chart_group_col": None,
        },
        "wordcount_by_gender": {
            "title": "Wordcount: By gender",
            "source_df": df,
            "stats_df": _calculate_grouped_wordcount_stats(
                df=df,
                group_cols=["gender"],
                bottom_percent=bottom_percent,
                logger=logger,
            ),
            "chart_group_col": "gender",
        },
        "wordcount_by_age_group": {
            "title": "Wordcount: By age group",
            "source_df": df,
            "stats_df": _calculate_grouped_wordcount_stats(
                df=df,
                group_cols=["age_group"],
                bottom_percent=bottom_percent,
                logger=logger,
            ),
            "chart_group_col": "age_group",
        },
    }

    for gender in ["female", "male", "other"]:
        gender_df = df.filter(pl.col("gender") == gender)

        if gender_df.is_empty():
            logger.warning(f"No rows found for gender: {gender}")
            continue

        sections[f"wordcount_by_age_group_for_{gender}"] = {
            "title": f"Wordcount: By age group for {gender}",
            "source_df": gender_df,
            "stats_df": _calculate_grouped_wordcount_stats(
                df=gender_df,
                group_cols=["age_group"],
                bottom_percent=bottom_percent,
                logger=logger,
            ),
            "chart_group_col": "age_group",
        }

    return sections


def _add_wordcount_interval(
    df: pl.DataFrame,
    wordcount_interval_size: int,
    wordcount_stop_at: int,
) -> pl.DataFrame:
    """Add wordcount interval labels to a dataframe."""
    return (
        df
        .with_columns(
            pl.when(pl.col("text_words") >= wordcount_stop_at)
            .then(wordcount_stop_at)
            .otherwise(
                (pl.col("text_words") // wordcount_interval_size)
                * wordcount_interval_size
            )
            .alias("interval_start")
        )
        .with_columns(
            pl.when(pl.col("interval_start") == wordcount_stop_at)
            .then(pl.lit(f"{wordcount_stop_at}+"))
            .otherwise(
                pl.col("interval_start").cast(pl.Utf8)
                + "-"
                + (
                    pl.col("interval_start") + wordcount_interval_size - 1
                ).cast(pl.Utf8)
            )
            .alias("wordcount_interval")
        )
    )


def _calculate_wordcount_distribution(
    df: pl.DataFrame,
    group_col: str | None,
    wordcount_interval_size: int,
    wordcount_stop_at: int,
    logger: logging.Logger,
) -> pl.DataFrame:
    """Calculate the distribution of texts across wordcount intervals."""
    logger.info(
        f"Calculating wordcount distribution. "
        f"group_col={group_col}, "
        f"wordcount_interval_size={wordcount_interval_size}, "
        f"wordcount_stop_at={wordcount_stop_at}"
    )

    df = df.filter(pl.col("text_words").is_not_null())

    if group_col is not None:
        df = df.filter(pl.col(group_col).is_not_null())

    df = _add_wordcount_interval(
        df=df,
        wordcount_interval_size=wordcount_interval_size,
        wordcount_stop_at=wordcount_stop_at,
    )

    if group_col is None:
        return (
            df
            .group_by(["interval_start", "wordcount_interval"])
            .agg(pl.len().alias("n_texts"))
            .sort("interval_start")
            .drop("interval_start")
        )

    return (
        df
        .group_by([group_col, "interval_start", "wordcount_interval"])
        .agg(pl.len().alias("n_texts"))
        .sort([group_col, "interval_start"])
        .drop("interval_start")
    )


def _save_wordcount_distribution_chart(
    distribution_df: pl.DataFrame,
    chart_title: str,
    chart_group_col: str | None,
    output_path: Path,
    logger: logging.Logger,
) -> None:
    """Save a wordcount distribution chart as a PNG image."""
    logger.info(f"Saving chart to {output_path}")

    plt.figure(figsize=(8, max(4, distribution_df.height * 0.18)))

    if chart_group_col is None:
        plt.barh(
            distribution_df["wordcount_interval"].to_list(),
            distribution_df["n_texts"].to_list(),
        )

    else:
        pivot_df = (
            distribution_df
            .pivot(
                values="n_texts",
                index="wordcount_interval",
                columns=chart_group_col,
                aggregate_function="sum",
            )
            .fill_null(0)
        )

        y_labels = pivot_df["wordcount_interval"].to_list()
        group_columns = [
            col for col in pivot_df.columns
            if col != "wordcount_interval"
        ]

        y_positions = list(range(len(y_labels)))
        bar_height = 0.8 / len(group_columns)

        for i, group_value in enumerate(group_columns):
            positions = [
                y + (i * bar_height) - (0.4 - bar_height / 2)
                for y in y_positions
            ]

            plt.barh(
                positions,
                pivot_df[group_value].to_list(),
                height=bar_height,
                label=str(group_value),
            )

        plt.yticks(y_positions, y_labels)
        plt.legend(title=chart_group_col.replace("_", " ").title())

    plt.title(chart_title)
    plt.xlabel("Number of texts")
    plt.ylabel("Wordcount interval")
    plt.tight_layout()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    plt.savefig(output_path, dpi=200)
    plt.close()


def _contact_reason_in_topic_expr(contact_reasons: list[str]) -> pl.Expr:
    """Create a Polars expression matching any contact reason in a topic."""
    return (
        pl.col("contact_reasons")
        .list.eval(pl.element().is_in(contact_reasons))
        .list.any()
    )


def _contact_reason_exists_expr(contact_reason: str) -> pl.Expr:
    """Create a Polars expression matching a specific contact reason."""
    return pl.col("contact_reasons").list.contains(contact_reason)


def _count_rows_for_label(
    df: pl.DataFrame,
    label: str,
    group_col: str | None,
    is_topic: bool,
) -> dict:
    """Count conversations for a topic or contact reason."""
    if group_col is None:
        return {
            "label": label,
            "is_topic": is_topic,
            "count": df.height,
        }

    counts = (
        df
        .filter(pl.col(group_col).is_not_null())
        .group_by(group_col)
        .agg(pl.len().alias("count"))
    )

    row = {
        "label": label,
        "is_topic": is_topic,
    }

    for result in counts.iter_rows(named=True):
        row[str(result[group_col])] = result["count"]

    return row


def _create_topic_contact_reason_pivot_table(
    df: pl.DataFrame,
    topic_contact_reasons: dict[str, list[str]],
    group_col: str | None,
    logger: logging.Logger,
) -> pl.DataFrame:
    """Create a pivot table of topics and contact reasons."""
    logger.info(f"Creating topic/contact reason pivot table. group_col={group_col}")

    df = df.filter(pl.col("contact_reasons").is_not_null())

    rows = []

    for topic, contact_reasons in topic_contact_reasons.items():
        topic_df = df.filter(_contact_reason_in_topic_expr(contact_reasons))

        rows.append(
            _count_rows_for_label(
                df=topic_df,
                label=topic,
                group_col=group_col,
                is_topic=True,
            )
        )

        for contact_reason in contact_reasons:
            contact_reason_df = df.filter(
                _contact_reason_exists_expr(contact_reason)
            )

            rows.append(
                _count_rows_for_label(
                    df=contact_reason_df,
                    label=f"  {contact_reason}",
                    group_col=group_col,
                    is_topic=False,
                )
            )

    table = pl.DataFrame(rows).fill_null(0)

    visible_count_cols = [
        col for col in table.columns
        if col not in ["label", "is_topic"]
    ]

    return table.select(["label", *visible_count_cols, "is_topic"])


def _create_topic_contact_reason_report_sections(
    df: pl.DataFrame,
    topic_contact_reasons: dict[str, list[str]],
    logger: logging.Logger,
) -> dict[str, dict]:
    """Create report sections for topic and contact-reason statistics."""
    logger.info("Creating topic/contact reason report sections.")

    sections = {
        "topics_entire_dataset": {
            "title": "Topics and contact reasons: Entire dataset",
            "stats_df": _create_topic_contact_reason_pivot_table(
                df=df,
                topic_contact_reasons=topic_contact_reasons,
                group_col=None,
                logger=logger,
            ),
        },
        "topics_by_gender": {
            "title": "Topics and contact reasons: By gender",
            "stats_df": _create_topic_contact_reason_pivot_table(
                df=df,
                topic_contact_reasons=topic_contact_reasons,
                group_col="gender",
                logger=logger,
            ),
        },
        "topics_by_age_group": {
            "title": "Topics and contact reasons: By age group",
            "stats_df": _create_topic_contact_reason_pivot_table(
                df=df,
                topic_contact_reasons=topic_contact_reasons,
                group_col="age_group",
                logger=logger,
            ),
        },
    }

    for gender in ["female", "male", "other"]:
        gender_df = df.filter(pl.col("gender") == gender)

        if gender_df.is_empty():
            logger.warning(f"No rows found for gender: {gender}")
            continue

        sections[f"topics_by_age_group_for_{gender}"] = {
            "title": f"Topics and contact reasons: By age group for {gender}",
            "stats_df": _create_topic_contact_reason_pivot_table(
                df=gender_df,
                topic_contact_reasons=topic_contact_reasons,
                group_col="age_group",
                logger=logger,
            ),
        }

    return sections


def _calculate_contact_reason_count_distribution(
    df: pl.DataFrame,
) -> pl.DataFrame:
    """Calculate the distribution of contact-reason counts per conversation."""
    return (
        df
        .filter(pl.col("contact_reasons").is_not_null())
        .with_columns(
            pl.col("contact_reasons")
            .list.len()
            .alias("n_contact_reasons")
        )
        .group_by("n_contact_reasons")
        .agg(pl.len().alias("count"))
        .sort("n_contact_reasons")
    )


def _calculate_topic_count_distribution(
    df: pl.DataFrame,
    topic_contact_reasons: dict[str, list[str]],
) -> pl.DataFrame:
    """Calculate the distribution of topic counts per conversation."""
    contact_reason_to_topic = {
        contact_reason: topic
        for topic, contact_reasons in topic_contact_reasons.items()
        for contact_reason in contact_reasons
    }

    rows = []

    for contact_reasons in (
        df
        .select("contact_reasons")
        .drop_nulls()
        .get_column("contact_reasons")
        .to_list()
    ):
        topics = {
            contact_reason_to_topic[contact_reason]
            for contact_reason in contact_reasons
            if contact_reason in contact_reason_to_topic
        }

        rows.append({"n_topics": len(topics)})

    return (
        pl.DataFrame(rows)
        .group_by("n_topics")
        .agg(pl.len().alias("count"))
        .sort("n_topics")
    )


def _calculate_topic_combination_distribution(
    df: pl.DataFrame,
    topic_contact_reasons: dict[str, list[str]],
) -> pl.DataFrame:
    """Calculate the frequency of topic combinations across conversations."""
    contact_reason_to_topic = {
        contact_reason: topic
        for topic, contact_reasons in topic_contact_reasons.items()
        for contact_reason in contact_reasons
    }

    rows = []

    for contact_reasons in (
        df
        .select("contact_reasons")
        .drop_nulls()
        .get_column("contact_reasons")
        .to_list()
    ):
        topics = sorted({
            contact_reason_to_topic[contact_reason]
            for contact_reason in contact_reasons
            if contact_reason in contact_reason_to_topic
        })

        rows.append({
            "n_topics": len(topics),
            "topic_combination": " + ".join(topics),
        })

    return (
        pl.DataFrame(rows)
        .group_by(["n_topics", "topic_combination"])
        .agg(pl.len().alias("count"))
        .sort("count", descending=True)
    )


def _calculate_wordcount_stats_by_number_of_contact_reasons(
    df: pl.DataFrame,
    bottom_percent: float,
    logger: logging.Logger,
) -> pl.DataFrame:
    """Calculate wordcount statistics by number of contact reasons."""
    logger.info(
        "Calculating wordcount stats grouped by number of contact reasons."
    )

    df_with_n_contact_reasons = (
        df
        .filter(pl.col("contact_reasons").is_not_null())
        .with_columns(
            pl.col("contact_reasons")
            .list.len()
            .alias("n_contact_reasons")
        )
        .select(
            "text_words",
            "n_contact_reasons",
        )
    )

    return (
        _calculate_grouped_wordcount_stats(
            df=df_with_n_contact_reasons,
            group_cols=["n_contact_reasons"],
            bottom_percent=bottom_percent,
            logger=logger,
        )
        .sort("n_contact_reasons")
    )


def _calculate_wordcount_stats_by_number_of_topics(
    df: pl.DataFrame,
    topic_contact_reasons: dict[str, list[str]],
    bottom_percent: float,
    logger: logging.Logger,
) -> pl.DataFrame:
    """Calculate wordcount statistics by number of topics."""
    logger.info("Calculating wordcount stats grouped by number of topics.")

    contact_reason_to_topic = {
        contact_reason: topic
        for topic, contact_reasons in topic_contact_reasons.items()
        for contact_reason in contact_reasons
    }

    rows = []

    for row in df.select(["text_words", "contact_reasons"]).iter_rows(named=True):
        contact_reasons = row["contact_reasons"]

        if contact_reasons is None:
            n_topics = 0
        else:
            n_topics = len({
                contact_reason_to_topic[contact_reason]
                for contact_reason in contact_reasons
                if contact_reason in contact_reason_to_topic
            })

        rows.append({
            "text_words": row["text_words"],
            "n_topics": n_topics,
        })

    df_with_n_topics = pl.DataFrame(rows)

    return (
        _calculate_grouped_wordcount_stats(
            df=df_with_n_topics,
            group_cols=["n_topics"],
            bottom_percent=bottom_percent,
            logger=logger,
        )
        .sort("n_topics")
    )


def _calculate_topic_frequency(
    df: pl.DataFrame,
    topic_contact_reasons: dict[str, list[str]],
) -> pl.DataFrame:
    """Calculate the number of conversations containing each topic."""
    rows = []

    for topic, contact_reasons in topic_contact_reasons.items():
        count = (
            df
            .filter(
                pl.col("contact_reasons")
                .list.eval(
                    pl.element().is_in(contact_reasons)
                )
                .list.any()
            )
            .height
        )

        rows.append({
            "topic": topic,
            "count": count,
        })

    return (
        pl.DataFrame(rows)
        .sort("count", descending=True)
    )


def _prettify_column_name(column_name: str) -> str:
    """Convert internal column names into readable report labels."""
    replacements = {
        "label": "Topic / contact reason",
        "group": "Group",
        "gender": "Gender",
        "age_group": "Age group",
        "n_texts": "Number of texts",
        "min_wordcount": "Minimum wordcount",
        "max_wordcount": "Maximum wordcount",
        "avg_wordcount": "Average wordcount",
        "med_wordcount": "Median wordcount",
        "count": "Count",
        "female": "Female",
        "male": "Male",
        "other": "Other",
        "pre teen": "Pre teen",
        "early teen": "Early teen",
        "teen": "Teen",
        "n_contact_reasons": "Number of contact reasons",
        "n_topics": "Number of topics",
        "topic_combination": "Topic combination",
        "topic": "Topic",
    }

    if column_name in replacements:
        return replacements[column_name]

    if column_name.startswith(
        "shortest_after_removing_bottom_"
    ):
        percent = (
            column_name
            .replace(
                "shortest_after_removing_bottom_",
                "",
            )
            .replace("_percent", "")
        )

        return (
            f"Shortest wordcount after removing "
            f"bottom {percent}%"
        )

    if column_name.startswith(
        "longest_after_removing_top_"
    ):
        percent = (
            column_name
            .replace(
                "longest_after_removing_top_",
                "",
            )
            .replace("_percent", "")
        )

        return (
            f"Longest wordcount after removing "
            f"top {percent}%"
        )

    return column_name.replace("_", " ").title()


def _set_cell_text(
    cell,
    text: str,
    bold: bool,
) -> None:
    """Set text and bold formatting for a Word table cell."""
    cell.text = ""

    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold


def _add_polars_table_to_docx(
    doc: Document,
    df: pl.DataFrame,
    title: str,
    logger: logging.Logger,
) -> None:
    """Add a Polars dataframe as a table in a Word document."""
    logger.info(f"Adding table to docx section: {title}")

    doc.add_heading(title, level=2)

    visible_columns = [
        col for col in df.columns
        if col != "is_topic"
    ]

    table = doc.add_table(rows=1, cols=len(visible_columns))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for i, column_name in enumerate(visible_columns):
        _set_cell_text(
            cell=header_cells[i],
            text=_prettify_column_name(column_name),
            bold=True,
        )

    for row in df.iter_rows(named=True):
        cells = table.add_row().cells
        is_topic = bool(row.get("is_topic", False))

        for i, column_name in enumerate(visible_columns):
            value = row[column_name]
            text = "" if value is None else str(value)

            _set_cell_text(
                cell=cells[i],
                text=text,
                bold=is_topic,
            )

    doc.add_paragraph()


def _add_chart_to_docx(
    doc: Document,
    chart_path: Path,
    logger: logging.Logger,
) -> None:
    """Add a chart image to a Word document."""
    logger.info(f"Adding chart to docx: {chart_path}")

    doc.add_picture(str(chart_path), width=Inches(5))
    doc.add_paragraph()


def _save_report_to_docx(
    wordcount_sections: dict[str, dict],
    topic_sections: dict[str, dict],
    contact_reason_count_distribution: pl.DataFrame,
    topic_count_distribution: pl.DataFrame,
    topic_combination_distribution: pl.DataFrame,
    topic_frequency: pl.DataFrame,
    wordcount_by_number_of_topics: pl.DataFrame,
    wordcount_by_number_of_contact_reasons: pl.DataFrame,
    output_dir_data: str,
    output_dir_pngs: str,
    wordcount_interval_size: int,
    wordcount_stop_at: int,
    logger: logging.Logger,
) -> None:
    """Build and save the full conversation data investigation report."""
    output_docx_path = Path(output_dir_data)
    output_docx_path.parent.mkdir(parents=True, exist_ok=True)

    output_png_dir = Path(output_dir_pngs)
    output_png_dir.mkdir(parents=True, exist_ok=True)

    logger.info(f"Creating docx report at {output_docx_path}")

    doc = Document()
    doc.add_heading("Conversation Data Investigation", level=1)

    doc.add_heading("Wordcount Statistics", level=1)

    for section_name, section in wordcount_sections.items():
        title = section["title"]
        source_df = section["source_df"]
        stats_df = section["stats_df"]
        chart_group_col = section["chart_group_col"]

        _add_polars_table_to_docx(
            doc=doc,
            df=stats_df,
            title=title,
            logger=logger,
        )

        distribution_df = _calculate_wordcount_distribution(
            df=source_df,
            group_col=chart_group_col,
            wordcount_interval_size=wordcount_interval_size,
            wordcount_stop_at=wordcount_stop_at,
            logger=logger,
        )

        chart_path = output_png_dir / f"{section_name}_wordcount_distribution.png"

        _save_wordcount_distribution_chart(
            distribution_df=distribution_df,
            chart_title=f"{title}: Wordcount distribution",
            chart_group_col=chart_group_col,
            output_path=chart_path,
            logger=logger,
        )

        _add_chart_to_docx(
            doc=doc,
            chart_path=chart_path,
            logger=logger,
        )

        doc.add_page_break()

    doc.add_heading("Topics Investigation", level=1)

    _add_polars_table_to_docx(
        doc=doc,
        df=topic_count_distribution,
        title="Number of topics per conversation",
        logger=logger,
    )

    _add_polars_table_to_docx(
        doc=doc,
        df=wordcount_by_number_of_topics,
        title="Wordcount statistics by number of topics",
        logger=logger,
    )

    _add_polars_table_to_docx(
        doc=doc,
        df=topic_frequency,
        title="Topic frequency",
        logger=logger,
    )

    doc.add_page_break()

    _add_polars_table_to_docx(
        doc=doc,
        df=contact_reason_count_distribution,
        title="Number of contact reasons per conversation",
        logger=logger,
    )

    _add_polars_table_to_docx(
        doc=doc,
        df=wordcount_by_number_of_contact_reasons,
        title="Wordcount statistics by number of contact reasons",
        logger=logger,
    )

    doc.add_page_break()

    _add_polars_table_to_docx(
        doc=doc,
        df=topic_combination_distribution,
        title="Topic combinations per conversation",
        logger=logger,
    )
    
    doc.add_page_break()

    doc.add_heading("Topics and Contact Reasons", level=1)

    for section in topic_sections.values():
        _add_polars_table_to_docx(
            doc=doc,
            df=section["stats_df"],
            title=section["title"],
            logger=logger,
        )

        doc.add_page_break()

    doc.save(output_docx_path)

    logger.info(f"Saved docx report to {output_docx_path}")


## MAIN FUNCTIONALITY ##
def investigate_data_stats(
    input_dir_data: str,
    output_dir_data: str,
    output_dir_pngs: str,
    topic_contact_reasons: dict[str, list[str]],
    logger: logging.Logger,
    bottom_percent: float,
    wordcount_interval_size: int,
    wordcount_stop_at: int,
) -> None:
    """Generate descriptive statistics and a report for ACT conversation data.
       Reads the cleaned conversation-level dataset, calculates wordcount,
           topic, and contact-reason statistics, generates wordcount
           distribution charts, and exports the results as a Word document.

       Args:
           input_dir_data: Path to the conversation-level parquet dataset.
           output_dir_data: Output path for the generated Word report.
           output_dir_pngs: Directory for intermediate chart images.
           topic_contact_reasons: Mapping of topics to their contact reasons.
           logger: Logger used to write processing information.
           bottom_percent: Percentage removed when calculating trimmed
               wordcount statistics.
           wordcount_interval_size: Width of each wordcount interval in
               distribution charts.
           wordcount_stop_at: Upper limit after which all wordcounts are
               grouped into a single interval.
    """
    logger.info(f"Reading input data from {input_dir_data}")

    df = pl.read_parquet(input_dir_data)

    logger.info(f"Loaded dataset with {df.height} rows and {df.width} columns.")

    wordcount_sections = _create_wordcount_report_sections(
        df=df,
        bottom_percent=bottom_percent,
        logger=logger,
    )

    topic_sections = _create_topic_contact_reason_report_sections(
        df=df,
        topic_contact_reasons=topic_contact_reasons,
        logger=logger,
    )

    contact_reason_count_distribution = _calculate_contact_reason_count_distribution(
        df=df,
    )

    topic_count_distribution = _calculate_topic_count_distribution(
        df=df,
        topic_contact_reasons=topic_contact_reasons,
    )

    topic_combination_distribution = _calculate_topic_combination_distribution(
        df=df,
        topic_contact_reasons=topic_contact_reasons,
    )

    wordcount_by_number_of_contact_reasons = (
        _calculate_wordcount_stats_by_number_of_contact_reasons(
            df=df,
            bottom_percent=bottom_percent,
            logger=logger,
        )
    )

    wordcount_by_number_of_topics = _calculate_wordcount_stats_by_number_of_topics(
        df=df,
        topic_contact_reasons=topic_contact_reasons,
        bottom_percent=bottom_percent,
        logger=logger,
    )

    topic_frequency = _calculate_topic_frequency(
        df=df,
        topic_contact_reasons=topic_contact_reasons,
    )

    _save_report_to_docx(
        wordcount_sections=wordcount_sections,
        topic_sections=topic_sections,
        contact_reason_count_distribution=contact_reason_count_distribution,
        topic_count_distribution=topic_count_distribution,
        topic_frequency=topic_frequency,
        topic_combination_distribution=topic_combination_distribution,
        wordcount_by_number_of_topics=wordcount_by_number_of_topics,
        wordcount_by_number_of_contact_reasons=wordcount_by_number_of_contact_reasons,
        output_dir_data=output_dir_data,
        output_dir_pngs=output_dir_pngs,
        wordcount_interval_size=wordcount_interval_size,
        wordcount_stop_at=wordcount_stop_at,
        logger=logger,
    )

    logger.info("Finished creating conversation data investigation report.")
## __________________ ##